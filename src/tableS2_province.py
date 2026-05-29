#!/usr/bin/env python3
"""
tableS2_province.py
===================
Supplementary Table S2 — Conservation coverage by province and instrument.

Reads pct_by_province.csv and writes a publication-ready formatted XLSX
(with headers grouped by level) plus a clean CSV in both English and
Spanish.

Outputs:
    TableS2_coverage_by_province.xlsx
    TableS2_coverage_by_province.csv
    TableS2_coverage_by_province_esp.xlsx
    TableS2_coverage_by_province_esp.csv

Authors
-------
    González Trilla G, Baldi G, Luchetti C, Pereira P & Grimson R (2025)
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parents[1]
PROC = ROOT / "data" / "processed"
OUT  = ROOT / "output" / "tables"
OUT.mkdir(parents=True, exist_ok=True)

# ── Categories ────────────────────────────────────────────────────────────────
CATS = ["C10", "C13", "C16", "C20", "C23", "C30", "C40", "C50"]

LEVEL_OF = {
    "C10": "A",
    "C13": "B", "C16": "B",
    "C20": "C", "C23": "C",
    "C30": "D", "C40": "D", "C50": "D",
}

# ── Colours per level (light tints for header backgrounds) ────────────────────
LEVEL_FILL = {
    "A": "C8E6C9",   # very light green
    "B": "DCEDC8",
    "C": "F0F4C3",
    "D": "FFF9C4",
}

# ── Localised text ────────────────────────────────────────────────────────────
LOCALES = {
    "en": {
        "title":          "Table S2 — Conservation coverage by province (% cumulative)",
        "col_province":   "Province",
        "col_area":       "Area (km²)",
        "col_total":      "Argentina (total)",
        "category_label": {
            "C10": "Strict PAs\n(IUCN I–IV)",
            "C13": "+ Glaciers\nLaw",
            "C16": "+ Forests\nCat. I",
            "C20": "+ Multi-use PAs\n(IUCN V–VI)",
            "C23": "+ Forests\nCat. II",
            "C30": "+ Unclassified\npublic PAs",
            "C40": "+ Private\nreserves",
            "C50": "+ Intl. & other",
        },
        "level_header":   lambda lv: f"Level {lv}",
        "province": {
            "Buenos Aires": "Buenos Aires", "CABA": "Buenos Aires City",
            "Catamarca": "Catamarca", "Chaco": "Chaco", "Chubut": "Chubut",
            "Córdoba": "Córdoba", "Cordoba": "Córdoba",
            "Corrientes": "Corrientes",
            "Entre Ríos": "Entre Ríos", "Entre Rios": "Entre Ríos",
            "Formosa": "Formosa", "Jujuy": "Jujuy",
            "La Pampa": "La Pampa", "La Rioja": "La Rioja",
            "Mendoza": "Mendoza", "Misiones": "Misiones",
            "Neuquén": "Neuquén", "Neuquen": "Neuquén",
            "Río Negro": "Río Negro", "Rio Negro": "Río Negro",
            "Salta": "Salta", "San Juan": "San Juan", "San Luis": "San Luis",
            "Santa Cruz": "Santa Cruz", "Santa Fe": "Santa Fe",
            "Santiago del Estero": "Santiago del Estero",
            "Tierra del Fuego": "Tierra del Fuego",
            "Tucumán": "Tucumán", "Tucuman": "Tucumán",
        },
        "suffix": "",
        "note": ("Values are cumulative percentages of provincial territory "
                 "protected at each level. Levels are nested: each column "
                 "adds the contribution of new instruments to all stricter "
                 "categories already included."),
    },
    "es": {
        "title":          "Tabla S2 — Cobertura de conservación por provincia (% acumulado)",
        "col_province":   "Provincia",
        "col_area":       "Superficie (km²)",
        "col_total":      "Argentina (total)",
        "category_label": {
            "C10": "ÁPs estrictas\n(UICN I–IV)",
            "C13": "+ Ley de\nGlaciares",
            "C16": "+ Bosques\nCat. I",
            "C20": "+ ÁPs uso múltiple\n(UICN V–VI)",
            "C23": "+ Bosques\nCat. II",
            "C30": "+ ÁPs públicas\nsin categoría",
            "C40": "+ Reservas\nprivadas",
            "C50": "+ Intl. y otras",
        },
        "level_header":   lambda lv: f"Nivel {lv}",
        "province": {
            "Buenos Aires": "Buenos Aires", "CABA": "CABA",
            "Catamarca": "Catamarca", "Chaco": "Chaco", "Chubut": "Chubut",
            "Córdoba": "Córdoba", "Cordoba": "Córdoba",
            "Corrientes": "Corrientes",
            "Entre Ríos": "Entre Ríos", "Entre Rios": "Entre Ríos",
            "Formosa": "Formosa", "Jujuy": "Jujuy",
            "La Pampa": "La Pampa", "La Rioja": "La Rioja",
            "Mendoza": "Mendoza", "Misiones": "Misiones",
            "Neuquén": "Neuquén", "Neuquen": "Neuquén",
            "Río Negro": "Río Negro", "Rio Negro": "Río Negro",
            "Salta": "Salta", "San Juan": "San Juan", "San Luis": "San Luis",
            "Santa Cruz": "Santa Cruz", "Santa Fe": "Santa Fe",
            "Santiago del Estero": "Santiago del Estero",
            "Tierra del Fuego": "Tierra del Fuego",
            "Tucumán": "Tucumán", "Tucuman": "Tucumán",
        },
        "suffix": "_esp",
        "note": ("Los valores son porcentajes acumulados del territorio "
                 "provincial protegido en cada nivel. Los niveles son "
                 "anidados: cada columna añade la contribución de "
                 "instrumentos nuevos a todas las categorías más estrictas "
                 "ya incluidas."),
    },
}


# ── Load and prepare data ─────────────────────────────────────────────────────
def load_data(locale: dict) -> tuple[pd.DataFrame, pd.Series]:
    """Returns (provinces dataframe sorted by total coverage, ARG total row)."""
    csv = PROC / "pct_by_province.csv"
    if not csv.exists():
        raise FileNotFoundError(
            f"Missing {csv}.\nRun 01_compute_coverage.py first."
        )
    df = pd.read_csv(csv)

    arg_row = df[df["cod_prov"] == "ARG"].iloc[0]

    df = df[df["cod_prov"] != "ARG"].copy()
    df["province_display"] = (
        df["province"].map(locale["province"]).fillna(df["province"])
    )
    df = df.sort_values("C50", ascending=False).reset_index(drop=True)
    return df, arg_row


# ── Clean CSV ─────────────────────────────────────────────────────────────────
def write_csv(df: pd.DataFrame, arg_row: pd.Series, locale: dict):
    """Write a tidy CSV with localised province names."""
    cols = ["province_display", "area_km2"] + CATS
    out = df[cols].copy()
    out.columns = [locale["col_province"], locale["col_area"]] + CATS

    arg = pd.Series({
        locale["col_province"]: locale["col_total"],
        locale["col_area"]:     arg_row["area_km2"],
        **{c: arg_row[c] for c in CATS},
    })
    out = pd.concat([out, pd.DataFrame([arg])], ignore_index=True)

    for c in CATS:
        out[c] = out[c].round(2)
    out[locale["col_area"]] = out[locale["col_area"]].round(0).astype(int)

    path = OUT / f"TableS2_coverage_by_province{locale['suffix']}.csv"
    out.to_csv(path, index=False, encoding="utf-8")
    print(f"Saved: {path}")


# ── Formatted XLSX ────────────────────────────────────────────────────────────
def write_xlsx(df: pd.DataFrame, arg_row: pd.Series, locale: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"

    # Styles
    thin   = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font  = Font(bold=True, size=10)
    body_font = Font(size=10)
    bold_font = Font(bold=True, size=10)

    # ── Row 1: title ──────────────────────────────────────────────────────────
    ws.cell(1, 1, locale["title"]).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=11)

    # ── Row 3-4: header (2 levels) ────────────────────────────────────────────
    LEVEL_ROW = 3
    CAT_ROW   = 4
    DATA_ROW  = 5

    # First two columns: province, area
    ws.cell(LEVEL_ROW, 1, locale["col_province"]).font = hdr_font
    ws.cell(LEVEL_ROW, 2, locale["col_area"]).font = hdr_font
    ws.merge_cells(start_row=LEVEL_ROW, end_row=CAT_ROW,
                   start_column=1, end_column=1)
    ws.merge_cells(start_row=LEVEL_ROW, end_row=CAT_ROW,
                   start_column=2, end_column=2)

    # Level header spans (A: 1 col, B: 2, C: 2, D: 3)
    col = 3
    for lv, cats_in_lv in [
        ("A", ["C10"]),
        ("B", ["C13", "C16"]),
        ("C", ["C20", "C23"]),
        ("D", ["C30", "C40", "C50"]),
    ]:
        span_start = col
        span_end   = col + len(cats_in_lv) - 1
        c = ws.cell(LEVEL_ROW, span_start, locale["level_header"](lv))
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center")
        c.fill = PatternFill("solid", fgColor=LEVEL_FILL[lv])
        if span_end > span_start:
            ws.merge_cells(start_row=LEVEL_ROW, end_row=LEVEL_ROW,
                           start_column=span_start, end_column=span_end)
        # Category sub-headers
        for cat in cats_in_lv:
            cell = ws.cell(CAT_ROW, col, locale["category_label"][cat])
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True,
                                       vertical="center")
            cell.fill = PatternFill("solid", fgColor=LEVEL_FILL[lv])
            col += 1

    # Apply borders to header
    for r in (LEVEL_ROW, CAT_ROW):
        for c in range(1, 11):
            ws.cell(r, c).border = border

    # ── Data rows ─────────────────────────────────────────────────────────────
    for i, (_, row) in enumerate(df.iterrows()):
        r = DATA_ROW + i
        ws.cell(r, 1, row["province_display"]).font = body_font
        ws.cell(r, 2, int(round(row["area_km2"]))).font = body_font
        ws.cell(r, 2).number_format = "#,##0"
        for j, cat in enumerate(CATS, start=3):
            cell = ws.cell(r, j, round(float(row[cat]), 2))
            cell.font = body_font
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="right")
        # Subtle banding for readability
        if i % 2 == 1:
            for c in range(1, 11):
                ws.cell(r, c).fill = PatternFill("solid", fgColor="F8F8F8")

    # ── National total row ────────────────────────────────────────────────────
    total_r = DATA_ROW + len(df)
    ws.cell(total_r, 1, locale["col_total"]).font = bold_font
    ws.cell(total_r, 2, int(round(arg_row["area_km2"]))).font = bold_font
    ws.cell(total_r, 2).number_format = "#,##0"
    for j, cat in enumerate(CATS, start=3):
        cell = ws.cell(total_r, j, round(float(arg_row[cat]), 2))
        cell.font = bold_font
        cell.number_format = "0.00"
        cell.alignment = Alignment(horizontal="right")
    for c in range(1, 11):
        ws.cell(total_r, c).fill = PatternFill("solid", fgColor="E8E8E8")
        ws.cell(total_r, c).border = Border(top=Side(style="medium",
                                                    color="555555"))

    # ── Note row ──────────────────────────────────────────────────────────────
    note_r = total_r + 2
    ws.cell(note_r, 1, locale["note"]).font = Font(italic=True, size=9,
                                                    color="555555")
    ws.merge_cells(start_row=note_r, end_row=note_r,
                   start_column=1, end_column=10)
    ws.cell(note_r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_r].height = 50

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    for letter in ["C", "D", "E", "F", "G", "H", "I", "J"]:
        ws.column_dimensions[letter].width = 14
    ws.row_dimensions[CAT_ROW].height = 40

    # Freeze panes
    ws.freeze_panes = "C5"

    path = OUT / f"TableS2_coverage_by_province{locale['suffix']}.xlsx"
    wb.save(path)
    print(f"Saved: {path}")


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    for lang in ("en", "es"):
        locale = LOCALES[lang]
        print(f"--- Building tables: {lang} ---")
        df, arg_row = load_data(locale)
        write_csv(df, arg_row, locale)
        write_xlsx(df, arg_row, locale)
    print("All done.")
