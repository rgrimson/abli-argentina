#!/usr/bin/env python3
"""
tableS1_ecoregion.py
====================
Supplementary Table S1 — Conservation coverage by ecoregion and instrument.

Reads pct_by_ecoregion.csv and writes a publication-ready formatted XLSX
(with headers grouped by level) plus a clean CSV in both English and
Spanish.

Outputs:
    TableS1_coverage_by_ecoregion.xlsx
    TableS1_coverage_by_ecoregion.csv
    TableS1_coverage_by_ecoregion_esp.xlsx
    TableS1_coverage_by_ecoregion_esp.csv

Authors
-------
    González Trilla G, Baldi G, Luchetti C, Pereira P & Grimson R (2025)
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parents[1]
PROC = ROOT / "data" / "processed"
OUT  = ROOT / "output" / "tables"
OUT.mkdir(parents=True, exist_ok=True)

# ── Categories ────────────────────────────────────────────────────────────────
CATS = ["C10", "C13", "C16", "C20", "C23", "C30", "C40", "C50"]

LEVEL_OF = {
    "C10": "A",
    "C13": "B", "C16": "B",
    "C20": "C", "C23": "C",
    "C30": "D", "C40": "D", "C50": "D",
}

LEVEL_FILL = {
    "A": "C8E6C9",
    "B": "DCEDC8",
    "C": "F0F4C3",
    "D": "FFF9C4",
}

# ── Localised text ────────────────────────────────────────────────────────────
LOCALES = {
    "en": {
        "title":          "Table S1 — Conservation coverage by ecoregion (% cumulative)",
        "col_ecoregion":  "Ecoregion",
        "col_area":       "Area (km²)",
        "category_label": {
            "C10": "Strict PAs\n(IUCN I–IV)",
            "C13": "+ Glaciers\nLaw",
            "C16": "+ Forests\nCat. I",
            "C20": "+ Multi-use PAs\n(IUCN V–VI)",
            "C23": "+ Forests\nCat. II",
            "C30": "+ Unclassified\npublic PAs",
            "C40": "+ Private\nreserves",
            "C50": "+ Intl. & other",
        },
        "level_header":   lambda lv: f"Level {lv}",
        "ecoregion": {
            "Altos Andes":                 "High Andes",
            "Bosques patagonicos":         "Patagonian Forests",
            "Chaco humedo":                "Humid Chaco",
            "Campos y Malezales":          "Campos and Malezales",
            "Chaco seco":                  "Dry Chaco",
            "Delta e islas del Parana":    "Paraná Delta and Islands",
            "Esteros del Ibera":           "Iberá Wetlands",
            "Estepa patagonica":           "Patagonian Steppe",
            "Espinal":                     "Espinal",
            "Monte de llanuras y mesetas": "Monte of Plains and Plateaus",
            "Monte de Sierras y Bolsones": "Monte of Sierras and Basins",
            "Pampa":                       "Pampas",
            "Puna":                        "Puna",
            "Selva Paranaense":            "Paranaense Forest",
            "Selva de Yungas":             "Yungas Forest",
        },
        "suffix": "",
        "note": ("Values are cumulative percentages of ecoregion territory "
                 "protected at each level. Levels are nested: each column "
                 "adds the contribution of new instruments to all stricter "
                 "categories already included."),
    },
    "es": {
        "title":          "Tabla S1 — Cobertura de conservación por ecorregión (% acumulado)",
        "col_ecoregion":  "Ecorregión",
        "col_area":       "Superficie (km²)",
        "category_label": {
            "C10": "ÁPs estrictas\n(UICN I–IV)",
            "C13": "+ Ley de\nGlaciares",
            "C16": "+ Bosques\nCat. I",
            "C20": "+ ÁPs uso múltiple\n(UICN V–VI)",
            "C23": "+ Bosques\nCat. II",
            "C30": "+ ÁPs públicas\nsin categoría",
            "C40": "+ Reservas\nprivadas",
            "C50": "+ Intl. y otras",
        },
        "level_header":   lambda lv: f"Nivel {lv}",
        "ecoregion": {
            "Altos Andes":                 "Altos Andes",
            "Bosques patagonicos":         "Bosques Patagónicos",
            "Chaco humedo":                "Chaco Húmedo",
            "Campos y Malezales":          "Campos y Malezales",
            "Chaco seco":                  "Chaco Seco",
            "Delta e islas del Parana":    "Delta e Islas del Paraná",
            "Esteros del Ibera":           "Esteros del Iberá",
            "Estepa patagonica":           "Estepa Patagónica",
            "Espinal":                     "Espinal",
            "Monte de llanuras y mesetas": "Monte de Llanuras y Mesetas",
            "Monte de Sierras y Bolsones": "Monte de Sierras y Bolsones",
            "Pampa":                       "Pampa",
            "Puna":                        "Puna",
            "Selva Paranaense":            "Selva Paranaense",
            "Selva de Yungas":             "Selva de Yungas",
        },
        "suffix": "_esp",
        "note": ("Los valores son porcentajes acumulados del territorio de "
                 "la ecorregión protegido en cada nivel. Los niveles son "
                 "anidados: cada columna añade la contribución de "
                 "instrumentos nuevos a todas las categorías más estrictas "
                 "ya incluidas."),
    },
}


# ── Load and prepare data ─────────────────────────────────────────────────────
def load_data(locale: dict) -> pd.DataFrame:
    csv = PROC / "pct_by_ecoregion.csv"
    if not csv.exists():
        raise FileNotFoundError(
            f"Missing {csv}.\nRun 01_compute_coverage.py first."
        )
    df = pd.read_csv(csv)
    df["ecoregion_display"] = (
        df["ecoregion"].map(locale["ecoregion"]).fillna(df["ecoregion"])
    )
    df = df.sort_values("C50", ascending=False).reset_index(drop=True)
    return df


# ── Clean CSV ─────────────────────────────────────────────────────────────────
def write_csv(df: pd.DataFrame, locale: dict):
    cols = ["ecoregion_display", "area_km2"] + CATS
    out = df[cols].copy()
    out.columns = [locale["col_ecoregion"], locale["col_area"]] + CATS

    for c in CATS:
        out[c] = out[c].round(2)
    out[locale["col_area"]] = out[locale["col_area"]].round(0).astype(int)

    path = OUT / f"TableS1_coverage_by_ecoregion{locale['suffix']}.csv"
    out.to_csv(path, index=False, encoding="utf-8")
    print(f"Saved: {path}")


# ── Formatted XLSX ────────────────────────────────────────────────────────────
def write_xlsx(df: pd.DataFrame, locale: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"

    thin   = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font  = Font(bold=True, size=10)
    body_font = Font(size=10)

    # Title
    ws.cell(1, 1, locale["title"]).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=11)

    LEVEL_ROW = 3
    CAT_ROW   = 4
    DATA_ROW  = 5

    ws.cell(LEVEL_ROW, 1, locale["col_ecoregion"]).font = hdr_font
    ws.cell(LEVEL_ROW, 2, locale["col_area"]).font = hdr_font
    ws.merge_cells(start_row=LEVEL_ROW, end_row=CAT_ROW,
                   start_column=1, end_column=1)
    ws.merge_cells(start_row=LEVEL_ROW, end_row=CAT_ROW,
                   start_column=2, end_column=2)

    col = 3
    for lv, cats_in_lv in [
        ("A", ["C10"]),
        ("B", ["C13", "C16"]),
        ("C", ["C20", "C23"]),
        ("D", ["C30", "C40", "C50"]),
    ]:
        span_start = col
        span_end   = col + len(cats_in_lv) - 1
        c = ws.cell(LEVEL_ROW, span_start, locale["level_header"](lv))
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center")
        c.fill = PatternFill("solid", fgColor=LEVEL_FILL[lv])
        if span_end > span_start:
            ws.merge_cells(start_row=LEVEL_ROW, end_row=LEVEL_ROW,
                           start_column=span_start, end_column=span_end)
        for cat in cats_in_lv:
            cell = ws.cell(CAT_ROW, col, locale["category_label"][cat])
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True,
                                       vertical="center")
            cell.fill = PatternFill("solid", fgColor=LEVEL_FILL[lv])
            col += 1

    for r in (LEVEL_ROW, CAT_ROW):
        for c in range(1, 11):
            ws.cell(r, c).border = border

    # Data rows
    for i, (_, row) in enumerate(df.iterrows()):
        r = DATA_ROW + i
        ws.cell(r, 1, row["ecoregion_display"]).font = body_font
        ws.cell(r, 2, int(round(row["area_km2"]))).font = body_font
        ws.cell(r, 2).number_format = "#,##0"
        for j, cat in enumerate(CATS, start=3):
            cell = ws.cell(r, j, round(float(row[cat]), 2))
            cell.font = body_font
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="right")
        if i % 2 == 1:
            for c in range(1, 11):
                ws.cell(r, c).fill = PatternFill("solid", fgColor="F8F8F8")

    # Note
    note_r = DATA_ROW + len(df) + 1
    ws.cell(note_r, 1, locale["note"]).font = Font(italic=True, size=9,
                                                    color="555555")
    ws.merge_cells(start_row=note_r, end_row=note_r,
                   start_column=1, end_column=10)
    ws.cell(note_r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_r].height = 50

    # Column widths
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    for letter in ["C", "D", "E", "F", "G", "H", "I", "J"]:
        ws.column_dimensions[letter].width = 14
    ws.row_dimensions[CAT_ROW].height = 40

    ws.freeze_panes = "C5"

    path = OUT / f"TableS1_coverage_by_ecoregion{locale['suffix']}.xlsx"
    wb.save(path)
    print(f"Saved: {path}")


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    for lang in ("en", "es"):
        locale = LOCALES[lang]
        print(f"--- Building tables: {lang} ---")
        df = load_data(locale)
        write_csv(df, locale)
        write_xlsx(df, locale)
    print("All done.")
